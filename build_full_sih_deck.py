import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -------------------------------------------------------------
# PRESENTATION CONSTANTS & COLOR PALETTE
# Serious Cybersecurity / AI Startup Aesthetic
# -------------------------------------------------------------
FONT_TITLE = "Helvetica"
FONT_BODY  = "Arial"

# Colors
C_BG_DARK       = RGBColor(10, 14, 23)      # #0A0E17 Deep Obsidian Navy
C_CARD_BG       = RGBColor(17, 28, 51)      # #111C33 Card Slate Navy
C_CARD_BORDER   = RGBColor(30, 46, 78)      # #1E2E4E Subtle Border
C_CARD_ACTIVE   = RGBColor(24, 40, 72)      # Highlighted Card Fill

C_CYAN          = RGBColor(0, 210, 255)     # #00D2FF Electric Cyan Accent
C_BLUE_ACCENT   = RGBColor(59, 130, 246)    # #3B82F6 Vibrant Cobalt
C_BLUE_DARK     = RGBColor(30, 58, 138)     # #1E3A8A Dark Blue

C_SPOOF_RED     = RGBColor(239, 68, 68)     # #EF4444 Alert Red
C_SPOOF_BG      = RGBColor(45, 20, 28)      # Red Tinted Dark Card
C_SPOOF_BORDER  = RGBColor(127, 29, 29)     # Dark Red Border

C_BONAFIDE_GREEN= RGBColor(16, 185, 129)    # #10B981 Verified Green
C_BONAFIDE_BG   = RGBColor(16, 38, 30)      # Green Tinted Dark Card
C_BONAFIDE_BORDER=RGBColor(6, 95, 70)       # Dark Green Border

C_AMBER         = RGBColor(245, 158, 11)    # #F59E0B Warning / Notice
C_AMBER_BG      = RGBColor(45, 35, 15)      # Amber Tinted Dark Card
C_AMBER_BORDER  = RGBColor(146, 64, 14)     # Dark Amber Border

C_TEXT_WHITE    = RGBColor(248, 250, 252)   # #F8FAFC Primary White
C_TEXT_MUTED    = RGBColor(148, 163, 184)   # #94A3B8 Slate Gray Body
C_TEXT_DIM      = RGBColor(100, 116, 139)   # #64748B Secondary Detail

SLIDE_WIDTH  = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

def create_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank_layout = prs.slide_layouts[6]
    return prs, blank_layout

def apply_dark_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_BG_DARK
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, category_text, slide_num):
    # Top Accent Line
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.35), Inches(11.733), Inches(0.04))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = C_CYAN
    top_line.line.fill.background()

    # Category Pill / Tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.48), Inches(7.5), Inches(0.3))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = f"SMART INDIA HACKATHON 2026  ·  {category_text.upper()}"
    p_tag.font.name = FONT_BODY
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = C_CYAN

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(9.8), Inches(0.7))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_t = tf_title.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(21)
    p_t.font.bold = True
    p_t.font.color.rgb = C_TEXT_WHITE

    # Badge on Right
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.633), Inches(0.48), Inches(1.9), Inches(0.52))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_CARD_BG
    badge.line.color.rgb = C_CARD_BORDER
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_b = tf_b.paragraphs[0]
    p_b.text = f"SIH-Aegis  |  {slide_num}/10"
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.name = FONT_BODY
    p_b.font.size = Pt(10.5)
    p_b.font.bold = True
    p_b.font.color.rgb = C_TEXT_MUTED

def set_speaker_notes(slide, meaning, script, qa_pairs):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.word_wrap = True
    
    full_text = []
    full_text.append("=== SLIDE MEANING & OBJECTIVE ===")
    full_text.append(meaning.strip())
    full_text.append("\n=== PRESENTER SCRIPT (20-40 SECONDS) ===")
    full_text.append(script.strip())
    full_text.append("\n=== LIKELY JUDGE QUESTIONS & DEFENSE ===")
    for q, a in qa_pairs:
        full_text.append(f"\nQ: {q}")
        full_text.append(f"A: {a}")
        
    tf.text = "\n".join(full_text)

def create_card(slide, left, top, width, height, border_color=C_CARD_BORDER, bg_color=C_CARD_BG):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

def add_bullet(tf, bold_prefix, text, size=10, bold_color=C_TEXT_WHITE, text_color=C_TEXT_MUTED):
    p = tf.add_paragraph()
    p.space_after = Pt(4)
    r1 = p.add_run()
    r1.text = bold_prefix + ": " if bold_prefix else ""
    r1.font.name = FONT_BODY
    r1.font.size = Pt(size)
    r1.font.bold = True
    r1.font.color.rgb = bold_color
    
    r2 = p.add_run()
    r2.text = text
    r2.font.name = FONT_BODY
    r2.font.size = Pt(size)
    r2.font.color.rgb = text_color

print("Base helper routines defined.")

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Color Palette ---
C_SIH_BLUE = RGBColor(0, 112, 192)       # #0070C0 (Official SIH Blue)
C_DARK_NAVY = RGBColor(11, 25, 44)       # #0B192C (Cyber Navy)
C_SLATE_BLUE = RGBColor(30, 58, 95)      # #1E3A5F (Header/Border Blue)
C_CYAN = RGBColor(0, 168, 232)           # #00A8E8 (High-Tech Cyan)
C_ACCENT_ORANGE = RGBColor(243, 156, 18) # #F39C12 (Amber Accent)
C_GREEN = RGBColor(39, 174, 96)          # #27AE60 (Bonafide Green)
C_RED = RGBColor(231, 76, 60)            # #E74C3C (Threat/Spoof Red)
C_DARK_TEXT = RGBColor(33, 37, 41)       # #212529 (Primary Dark Text)
C_MUTED_TEXT = RGBColor(100, 110, 120)   # #646E78 (Secondary Text)
C_WHITE = RGBColor(255, 255, 255)
C_CARD_BG = RGBColor(248, 250, 252)      # #F8FAFC (Clean Light Card)
C_CARD_BORDER = RGBColor(218, 225, 233)  # #DAE1E9 (Subtle Border)
C_DARK_CARD_BG = RGBColor(15, 23, 42)    # #0F172A (Cyber Dark Card)
C_LIGHT_BLUE_BG = RGBColor(240, 247, 255) # Light Blue Tint
C_LIGHT_RED_BG = RGBColor(254, 242, 242) # Light Red Tint
C_LIGHT_GREEN_BG = RGBColor(240, 253, 244) # Light Green Tint

FONT_TITLE = "Times New Roman"
FONT_BODY = "Arial"

def add_slide_chrome(slide, title_text, subtitle_text, slide_num, logo_path='sih_logo.png'):
    # 1. Oval (Top-Left)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, 329773, 252246, 1251857, 807334)
    oval.fill.background()
    oval.line.color.rgb = RGBColor(90, 95, 100)
    oval.line.width = Pt(1.5)
    tf = oval.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "Your\nTeam\nName"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_TITLE
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(50, 50, 50)
    
    # 2. SIH 2026 Logo (Top-Right)
    if os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, 9780086, 1500, 2249850, 1062337)
        
    # 3. Bottom Blue Banner
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 6354762, 12192000, 503238)
    banner.fill.solid()
    banner.fill.fore_color.rgb = C_SIH_BLUE
    banner.line.fill.background()
    
    # 4. Footer Text
    fb = slide.shapes.add_textbox(3600000, 6380000, 5000000, 400000)
    tf_f = fb.text_frame
    tf_f.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_f = tf_f.paragraphs[0]
    p_f.text = "@SIH Idea submission- Template"
    p_f.alignment = PP_ALIGN.CENTER
    p_f.font.name = FONT_BODY
    p_f.font.size = Pt(12)
    p_f.font.color.rgb = C_WHITE
    p_f.font.bold = False
    
    # 5. Slide Number
    nb = slide.shapes.add_textbox(10800000, 6380000, 1100000, 400000)
    tf_n = nb.text_frame
    tf_n.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_n = tf_n.paragraphs[0]
    p_n.text = str(slide_num)
    p_n.alignment = PP_ALIGN.RIGHT
    p_n.font.name = FONT_BODY
    p_n.font.size = Pt(12)
    p_n.font.color.rgb = C_WHITE
    p_n.font.bold = True

    # 6. Title Box
    tb = slide.shapes.add_textbox(1700000, 90000, 7900000, 580000)
    tf_t = tb.text_frame
    tf_t.word_wrap = True
    tf_t.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text.upper()
    p_t.alignment = PP_ALIGN.CENTER
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = C_DARK_NAVY
    
    # 7. Subtitle Box
    if subtitle_text:
        sb = slide.shapes.add_textbox(1700000, 660000, 7900000, 400000)
        tf_s = sb.text_frame
        tf_s.word_wrap = True
        tf_s.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.alignment = PP_ALIGN.CENTER
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(13)
        p_s.font.bold = True
        p_s.font.color.rgb = C_SIH_BLUE

def add_card(slide, left, top, width, height, title, items, 
             bg_color=C_CARD_BG, border_color=C_CARD_BORDER, 
             accent_color=None, title_color=C_DARK_NAVY, font_size=10.5, title_size=12.5):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
        
    if accent_color:
        strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, int(Pt(5)))
        strip.fill.solid()
        strip.fill.fore_color.rgb = accent_color
        strip.line.fill.background()

    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = int(Pt(12))
    tf.margin_right = int(Pt(12))
    tf.margin_top = int(Pt(10))
    tf.margin_bottom = int(Pt(10))

    if title:
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = FONT_BODY
        p_title.font.size = Pt(title_size)
        p_title.font.bold = True
        p_title.font.color.rgb = title_color
        p_title.space_after = Pt(5)
        first = True
    else:
        first = False

    for item in items:
        if first and not title:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        
        if ":" in item and len(item.split(":", 1)[0]) < 40:
            prefix, rest = item.split(":", 1)
            r1 = p.add_run()
            r1.text = "• " + prefix + ":"
            r1.font.name = FONT_BODY
            r1.font.size = Pt(font_size)
            r1.font.bold = True
            r1.font.color.rgb = title_color if bg_color == C_DARK_CARD_BG else C_DARK_TEXT
            
            r2 = p.add_run()
            r2.text = rest
            r2.font.name = FONT_BODY
            r2.font.size = Pt(font_size)
            r2.font.bold = False
            r2.font.color.rgb = C_WHITE if bg_color == C_DARK_CARD_BG else C_DARK_TEXT
        else:
            p.text = "• " + item if not item.startswith("• ") and not item.startswith("  ") else item
            p.font.name = FONT_BODY
            p.font.size = Pt(font_size)
            p.font.color.rgb = C_WHITE if bg_color == C_DARK_CARD_BG else C_DARK_TEXT
        p.space_after = Pt(3.5)
    return card

def add_badge(slide, left, top, width, height, text, bg_color=C_SIH_BLUE, text_color=C_WHITE, font_size=9.5):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = bg_color
    box.line.fill.background()
    tf = box.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_BODY
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = text_color
    return box

def set_notes(slide, text):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = text

print("build_full_presentation base ready.")

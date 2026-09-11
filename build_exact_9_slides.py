import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Official SIH Palette ---
C_SIH_BLUE = RGBColor(0, 112, 192)         # #0070C0 Official SIH Blue
C_NAVY_DARK = RGBColor(11, 37, 69)         # #0B2545 Deep SIH Navy
C_LIGHT_BLUE_BG = RGBColor(235, 245, 255)  # Flow / Header soft blue
C_BORDER_GRAY = RGBColor(218, 224, 233)    # Subtle card borders
C_CARD_BG = RGBColor(248, 250, 252)        # Off-white card fill
C_TEXT_DARK = RGBColor(20, 24, 30)         # Body text
C_TEXT_MUTED = RGBColor(90, 100, 110)      # Secondary text
C_WHITE = RGBColor(255, 255, 255)
C_ACCENT_ORANGE = RGBColor(230, 126, 34)   # #E67E22
C_ACCENT_GREEN = RGBColor(39, 174, 96)     # #27AE60
C_ACCENT_RED = RGBColor(231, 76, 60)       # #E74C3C
C_FLOW_BOX_BG = RGBColor(238, 245, 253)    # Flowchart step background
C_FLOW_BOX_BORDER = RGBColor(180, 210, 240)

FONT_TITLE = "Times New Roman"
FONT_BODY = "Arial"

def add_chrome(slide, title_text, subtitle_text, slide_num, logo_path='sih_logo.png'):
    # 1. Top-Left Team Name Oval
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, 329773, 252246, 1251857, 807334)
    oval.fill.background()
    oval.line.color.rgb = RGBColor(80, 80, 80)
    oval.line.width = Pt(1.5)
    tf_o = oval.text_frame
    tf_o.word_wrap = True
    tf_o.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_o = tf_o.paragraphs[0]
    p_o.text = "SIH-Aegis"
    p_o.alignment = PP_ALIGN.CENTER
    p_o.font.name = FONT_TITLE
    p_o.font.size = Pt(14)
    p_o.font.bold = True
    p_o.font.color.rgb = RGBColor(30, 30, 30)

    # 2. Top-Right SIH 2026 Logo
    if os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, 9780086, 1500, 2249850, 1062337)

    # 3. Bottom Blue Banner (#0070C0)
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 6354762, 12192000, 503238)
    banner.fill.solid()
    banner.fill.fore_color.rgb = C_SIH_BLUE
    banner.line.fill.background()

    # 4. Footer Placeholder Text
    fb = slide.shapes.add_textbox(3600000, 6380000, 5000000, 400000)
    tf_f = fb.text_frame
    tf_f.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_f = tf_f.paragraphs[0]
    p_f.text = "@SIH Idea submission- Template"
    p_f.alignment = PP_ALIGN.CENTER
    p_f.font.name = FONT_BODY
    p_f.font.size = Pt(12)
    p_f.font.color.rgb = C_WHITE

    # 5. Slide Number
    nb = slide.shapes.add_textbox(10800000, 6380000, 1100000, 400000)
    tf_n = nb.text_frame
    tf_n.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_n = tf_n.paragraphs[0]
    p_n.text = str(slide_num)
    p_n.alignment = PP_ALIGN.RIGHT
    p_n.font.name = FONT_BODY
    p_n.font.size = Pt(12)
    p_n.font.color.rgb = C_WHITE
    p_n.font.bold = True

    # 6. Title Box
    tb = slide.shapes.add_textbox(1650000, 90000, 8000000, 560000)
    tf_t = tb.text_frame
    tf_t.word_wrap = True
    tf_t.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text.upper()
    p_t.alignment = PP_ALIGN.CENTER
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = RGBColor(20, 20, 20)

    # 7. Subtitle Box
    if subtitle_text:
        sb = slide.shapes.add_textbox(1650000, 660000, 8000000, 420000)
        tf_s = sb.text_frame
        tf_s.word_wrap = True
        tf_s.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.alignment = PP_ALIGN.CENTER
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(13)
        p_s.font.bold = True
        p_s.font.color.rgb = C_SIH_BLUE

def add_flow_step(slide, left, top, width, height, title, subtitle, bg_color=C_FLOW_BOX_BG, border_color=C_FLOW_BOX_BORDER, text_color=C_NAVY_DARK):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = bg_color
    box.line.color.rgb = border_color
    box.line.width = Pt(1)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = int(Pt(6))
    tf.margin_right = int(Pt(6))
    tf.margin_top = int(Pt(4))
    tf.margin_bottom = int(Pt(4))
    
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = FONT_BODY
    p1.font.size = Pt(10)
    p1.font.bold = True
    p1.font.color.rgb = text_color
    
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.alignment = PP_ALIGN.CENTER
        p2.font.name = FONT_BODY
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = C_TEXT_MUTED if text_color != C_WHITE else C_WHITE

def add_arrow(slide, left, top, width=220000, height=180000):
    ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    ar.fill.solid()
    ar.fill.fore_color.rgb = C_SIH_BLUE
    ar.line.fill.background()

def add_column_card(slide, left, top, width, height, title, items, font_size=9.5, title_size=11.5, bullet_color=C_TEXT_DARK):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD_BG
    card.line.color.rgb = C_BORDER_GRAY
    card.line.width = Pt(1)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = int(Pt(12))
    tf.margin_right = int(Pt(12))
    tf.margin_top = int(Pt(12))
    tf.margin_bottom = int(Pt(10))
    
    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.name = FONT_BODY
    p0.font.size = Pt(title_size)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY_DARK
    p0.space_after = Pt(6)
    
    for item in items:
        p = tf.add_paragraph()
        p.text = "• " + item if not item.startswith("• ") and not item.startswith("  ") else item
        p.font.name = FONT_BODY
        p.font.size = Pt(font_size)
        p.font.color.rgb = bullet_color
        p.space_after = Pt(4)
    return card

def set_notes(slide, text):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = text

print("build_exact_9_slides base ready.")
